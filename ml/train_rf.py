"""
Train and evaluate the Random Forest for the IT Career Path Recommender.

Reads Job_Posting_Review_Workbook.xlsx (Review + Skill Matrix + Career Skill Profiles sheets),
uses rows with "In dataset? = Yes" and the Final label as the target.

Usage:
    python train_rf.py --workbook Job_Posting_Review_Workbook.xlsx --features skills
    python train_rf.py --workbook Job_Posting_Review_Workbook.xlsx --features domain

--features skills : one 0/1 column per skill (284 features)
--features domain : the MIX of the six career areas in a posting: each area's share of the posting's
                    detected area skills (6 features that sum to 1). A student's tiered assessment is
                    mapped into this same representation (see student_mapping.py).
--final           : tune on all rows, fit on all rows, and save final_model.joblib for the backend.

Open and save the workbook in Excel/LibreOffice first so its formulas have cached values.
"""
import argparse, json, re, warnings
import numpy as np, pandas as pd
from openpyxl import load_workbook
from sklearn.ensemble import RandomForestClassifier
from sklearn.tree import DecisionTreeClassifier
from sklearn.model_selection import StratifiedGroupKFold, RandomizedSearchCV, GridSearchCV
from sklearn.metrics import (accuracy_score, f1_score, classification_report, confusion_matrix)
from sklearn.calibration import CalibratedClassifierCV
import joblib
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt

warnings.filterwarnings("ignore")
SEED = 42

def load(path):
    wb = load_workbook(path, data_only=True)
    rv = wb["Review"]
    hdr = [c.value for c in rv[1]]; ix = {h: i for i, h in enumerate(hdr)}
    rev = pd.DataFrame([{
        "rid": r[ix["Row ID"]], "title": r[ix["Job title"]], "company": r[ix["Company"]],
        "source": r[ix["Source"]], "final": r[ix["Final label"]], "in": r[ix["In dataset?"]]}
        for r in rv.iter_rows(min_row=2, values_only=True) if r[0]])
    if rev["in"].isna().all():
        raise SystemExit("Formula values are empty. Open the workbook in Excel/LibreOffice, save it, and rerun.")
    sk = wb["Skill Matrix"]
    # Find the header row by its "Row ID" cell instead of assuming a fixed row number,
    # so the script still works if the note rows above the table are deleted or added.
    hrow = next((i for i, r in enumerate(sk.iter_rows(min_row=1, max_row=15, values_only=True), 1)
                 if r and r[0] == "Row ID"), None)
    if hrow is None:
        raise SystemExit("Could not find the 'Row ID' header row in the Skill Matrix sheet.")
    heads = [c.value for c in sk[hrow]]
    start = heads.index("Skills found") + 1          # skills begin right after this column
    skills = [h for h in heads[start:] if h not in (None, "")]
    if len(skills) != len(set(skills)):
        raise SystemExit("Duplicate skill column names in the Skill Matrix header.")
    rows = [[r[0]] + [1 if v == 1 else 0 for v in r[start:start + len(skills)]]
            for r in sk.iter_rows(min_row=hrow + 1, values_only=True)
            if r[0] and str(r[0]).startswith("R")]
    mat = pd.DataFrame(rows, columns=["rid"] + skills)
    areas = {r[0]: r[1] for r in wb["Career Skill Profiles"].iter_rows(min_row=6, values_only=True) if r[0]}
    return rev, mat, skills, areas

CAREER_AREAS = ["Software development", "Networking", "Cybersecurity",
                "Data science & analytics", "Database administration", "UI/UX & web development"]

def domain_features(X, areas):
    """Share of each career area in the posting's detected skills (rows sum to 1; all-zero rows stay 0).
    The 'General / shared' area is left out because a student assessment does not measure it."""
    counts = pd.DataFrame({a: X[[c for c in X.columns if areas.get(c) == a]].sum(axis=1) for a in CAREER_AREAS})
    tot = counts.sum(axis=1)
    return counts.div(tot.where(tot > 0, 1), axis=0)

def topk_accuracy(proba, classes, y_true, k=3):
    top = np.argsort(-proba, axis=1)[:, :k]
    return float(np.mean([y_true[i] in set(np.array(classes)[top[i]]) for i in range(len(y_true))]))

def company_key(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default="Job_Posting_Review_Workbook.xlsx")
    ap.add_argument("--features", choices=["skills", "domain"], default="skills")
    ap.add_argument("--outdir", default="model_outputs")
    ap.add_argument("--final", action="store_true", help="tune on all rows, fit on all rows, save final_model.joblib")
    ap.add_argument("--calibrate", action="store_true", help="wrap the forest in sigmoid calibration")
    a = ap.parse_args()
    import os; os.makedirs(a.outdir, exist_ok=True)

    rev, mat, skills, areas = load(a.workbook)
    df = rev[rev["in"] == "Yes"].merge(mat, on="rid").reset_index(drop=True)
    y = df["final"].astype(str).values
    groups = df["company"].map(company_key).values
    Xs = df[skills]
    X = Xs if a.features == "skills" else domain_features(Xs, areas)
    classes = sorted(set(y))
    log = []
    def say(s=""):
        print(s); log.append(s)
    say(f"Rows used: {len(df)} | features: {a.features} ({X.shape[1]}) | companies: {len(set(groups))}")
    say(pd.Series(y).value_counts().to_string()); say()
    for c, n in pd.Series(y).value_counts().items():
        if n < 50: say(f"WARNING: class '{c}' has only {n} rows.")

    # Grouped, stratified split: postings from the same company never appear in both train and test.
    sgkf = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=SEED)
    splits = list(sgkf.split(X, y, groups))
    tr, te = splits[0]
    Xtr, Xte, ytr, yte = X.iloc[tr], X.iloc[te], y[tr], y[te]
    say(f"Holdout split: train {len(tr)} / test {len(te)} (grouped by company)")

    inner = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=SEED)
    inner_cv = list(inner.split(Xtr, ytr, groups[tr]))

    # Baseline: single decision tree
    dt = GridSearchCV(DecisionTreeClassifier(class_weight="balanced", random_state=SEED),
                      {"max_depth": [3, 5, 8, 12, None], "min_samples_leaf": [1, 3, 5]},
                      scoring="f1_macro", cv=inner_cv).fit(Xtr, ytr)
    # Random Forest
    rf = RandomizedSearchCV(
        RandomForestClassifier(class_weight="balanced", random_state=SEED, n_jobs=-1),
        {"n_estimators": [200, 300, 500], "max_depth": [None, 8, 12, 20],
         "max_features": ["sqrt", "log2", 0.3], "min_samples_leaf": [1, 2, 4]},
        n_iter=20, scoring="f1_macro", cv=inner_cv, random_state=SEED, n_jobs=-1).fit(Xtr, ytr)
    best = rf.best_estimator_
    say(f"Best RF params: {rf.best_params_}")
    model = CalibratedClassifierCV(best, method="sigmoid", cv=5).fit(Xtr, ytr) if a.calibrate else best

    res = {}
    for name, m in [("Decision Tree (baseline)", dt.best_estimator_), ("Random Forest", model)]:
        pred = m.predict(Xte); proba = m.predict_proba(Xte)
        res[name] = {"accuracy": accuracy_score(yte, pred), "macro_f1": f1_score(yte, pred, average="macro"),
                     "top3_accuracy": topk_accuracy(proba, list(m.classes_), yte)}
    say("\nHoldout test results"); say(pd.DataFrame(res).T.round(3).to_string())
    pred = model.predict(Xte)
    say("\nRandom Forest per-class report"); say(classification_report(yte, pred, digits=3))
    cm = confusion_matrix(yte, pred, labels=classes)
    say("Confusion matrix (rows = true, columns = predicted)")
    say(pd.DataFrame(cm, index=classes, columns=classes).to_string())

    # Grouped 5-fold CV with the chosen settings (more stable than one split on a small dataset)
    oof = np.empty(len(y), dtype=object); oofp = np.zeros((len(y), len(classes)))
    for trn, tst in splits:
        m = RandomForestClassifier(**{**best.get_params()}).fit(X.iloc[trn], y[trn])
        oof[tst] = m.predict(X.iloc[tst]); oofp[tst] = m.predict_proba(X.iloc[tst])
    f1s = [f1_score(y[tst], oof[tst], average="macro") for _, tst in splits]
    say(f"\nGrouped 5-fold CV: accuracy {accuracy_score(y, oof):.3f} | macro-F1 {np.mean(f1s):.3f} +/- {np.std(f1s):.3f} | "
        f"top-3 {topk_accuracy(oofp, classes, y):.3f}")

    # Source-bias check: train on Greenhouse/Lever, test on Himalayas (Networking and Database only)
    him = (df["source"] == "Himalayas").values
    if him.sum() and (~him).sum():
        m = RandomForestClassifier(**best.get_params()).fit(X[~him], y[~him])
        p = m.predict(X[him])
        say(f"\nSource check: trained on Greenhouse/Lever ({(~him).sum()} rows), tested on Himalayas ({him.sum()} rows): "
            f"accuracy {accuracy_score(y[him], p):.3f}")
        say("If this is much lower than the holdout accuracy, the model may be learning posting style, not job content.")

    # Feature importance
    imp = pd.Series(best.feature_importances_, index=X.columns).sort_values(ascending=False)
    imp.head(40).rename("importance").to_csv(f"{a.outdir}/feature_importance_{a.features}.csv")
    say("\nTop 15 features by importance"); say(imp.head(15).round(4).to_string())

    if a.final:
        import sklearn, datetime, hashlib
        cvall = list(StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=SEED).split(X, y, groups))
        search = RandomizedSearchCV(
            RandomForestClassifier(class_weight="balanced", random_state=SEED, n_jobs=-1),
            {"n_estimators": [200, 300, 500], "max_depth": [None, 8, 12, 20],
             "max_features": ["sqrt", "log2", 0.3, 0.5], "min_samples_leaf": [1, 2, 4]},
            n_iter=20, scoring="f1_macro", cv=cvall, random_state=SEED, n_jobs=-1).fit(X, y)
        fin = search.best_estimator_
        say(f"\nFINAL model: best grouped-CV macro-F1 {search.best_score_:.3f} with {search.best_params_}")
        sk_mat = df[skills]
        profiles = {c: sk_mat[y == c].mean().round(4).to_dict() for c in classes}          # skill share by category
        mix_median = {c: X[y == c].median().round(4).to_dict() for c in classes}            # typical area mix per category
        bundle = {"model": fin, "features": list(X.columns), "classes": list(fin.classes_), "mode": a.features,
                  "areas": areas, "skill_profiles": profiles, "mix_median": mix_median,
                  "meta": {"trained_rows": int(len(y)), "class_counts": pd.Series(y).value_counts().to_dict(),
                           "cv_macro_f1": float(search.best_score_), "params": search.best_params_,
                           "sklearn": sklearn.__version__, "trained_at": datetime.datetime.now().isoformat(timespec="seconds"),
                           "workbook_sha1": hashlib.sha1(open(a.workbook, "rb").read()).hexdigest()[:12],
                           "label_status": "DRAFT labels unless the team has verified them in the workbook"}}
        joblib.dump(bundle, f"{a.outdir}/final_model.joblib")
        say(f"Saved {a.outdir}/final_model.joblib (trained on all {len(y)} rows)")

    # Outputs
    fig, ax = plt.subplots(figsize=(8, 6.5))
    ax.imshow(cm, cmap="Blues")
    ax.set_xticks(range(len(classes))); ax.set_yticks(range(len(classes)))
    short = [c.replace(" & ", "/").replace("development", "dev").replace("administration", "admin") for c in classes]
    ax.set_xticklabels(short, rotation=40, ha="right"); ax.set_yticklabels(short)
    for i in range(len(classes)):
        for j in range(len(classes)):
            ax.text(j, i, cm[i, j], ha="center", va="center", color="white" if cm[i, j] > cm.max() / 2 else "black")
    ax.set_xlabel("Predicted"); ax.set_ylabel("True"); ax.set_title(f"Random Forest holdout confusion matrix ({a.features})")
    plt.tight_layout(); plt.savefig(f"{a.outdir}/confusion_matrix_{a.features}.png", dpi=150); plt.close()
    joblib.dump({"model": model, "features": list(X.columns), "classes": list(model.classes_),
                 "mode": a.features, "areas": areas}, f"{a.outdir}/rf_model_{a.features}.joblib")
    pd.DataFrame(res).T.to_csv(f"{a.outdir}/holdout_metrics_{a.features}.csv")
    open(f"{a.outdir}/results_{a.features}.txt", "w").write("\n".join(log))
    say(f"\nSaved outputs to {a.outdir}/")

if __name__ == "__main__":
    main()